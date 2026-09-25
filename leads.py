"""Modelo de lead, deduplicação, classificação e exportação (CSV/XLSX)."""

import csv
import re
import unicodedata
from dataclasses import asdict, dataclass, fields
from urllib.parse import urlparse

REDES_SOCIAIS = (
    "instagram.com", "facebook.com", "fb.com", "fb.me", "linktr.ee",
    "wa.me", "whatsapp.com", "tiktok.com", "bio.link", "linkbio",
    "beacons.ai", "ifood.com.br", "goomer", "anota.ai", "sites.google.com",
)


@dataclass
class Lead:
    nome: str
    categoria: str = ""
    endereco: str = ""
    telefone: str = ""
    site: str = ""
    nota: float | None = None
    avaliacoes: int = 0
    status: str = ""
    latitude: float | None = None
    longitude: float | None = None
    link_maps: str = ""
    place_id: str = ""
    busca: str = ""


def _normalizar(texto: str) -> str:
    texto = unicodedata.normalize("NFKD", texto or "").encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]", "", texto.lower())


def so_digitos(texto: str) -> str:
    return re.sub(r"\D", "", texto or "")


def chave(lead: Lead) -> str:
    if lead.place_id:
        return lead.place_id
    return _normalizar(lead.nome) + "|" + (so_digitos(lead.telefone) or _normalizar(lead.endereco))


def situacao_site(site: str) -> str:
    if not site:
        return "SEM SITE"
    host = urlparse(site if "://" in site else "http://" + site).netloc.lower()
    if any(r in host or r in site.lower() for r in REDES_SOCIAIS):
        return "SÓ REDE SOCIAL"
    return "TEM SITE"


def whatsapp(telefone: str) -> str:
    """Link wa.me se o número parece celular (DDD + 9 dígitos começando em 9)."""
    d = so_digitos(telefone)
    if d.startswith("55") and len(d) in (12, 13):
        d = d[2:]
    if len(d) == 11 and d[2] == "9":
        return f"https://wa.me/55{d}"
    return ""


def prioridade(lead: Lead) -> str:
    sit = situacao_site(lead.site)
    if sit == "TEM SITE":
        return "C"
    if sit == "SEM SITE" and lead.telefone:
        return "A"
    return "B"


def na_cidade(lead: Lead, cidade: str) -> bool:
    return _normalizar(cidade) in _normalizar(lead.endereco)


class Colecao:
    """Guarda leads sem duplicar, mesclando dados quando o mesmo local reaparece."""

    def __init__(self):
        self._itens: dict[str, Lead] = {}

    def __len__(self):
        return len(self._itens)

    def __contains__(self, k: str):
        return k in self._itens

    def adicionar(self, lead: Lead) -> bool:
        k = chave(lead)
        existente = self._itens.get(k)
        if existente is None:
            self._itens[k] = lead
            return True
        for f in fields(Lead):
            if not getattr(existente, f.name) and getattr(lead, f.name):
                setattr(existente, f.name, getattr(lead, f.name))
        if lead.busca and lead.busca not in existente.busca.split(", "):
            existente.busca = f"{existente.busca}, {lead.busca}".strip(", ")
        return False

    def ordenados(self) -> list[Lead]:
        return sorted(self._itens.values(), key=lambda l: (prioridade(l), -(l.avaliacoes or 0), l.nome))


COLUNAS = [
    "prioridade", "situacao_site", "nome", "categoria", "telefone", "whatsapp",
    "site", "endereco", "nota", "avaliacoes", "status", "link_maps",
    "latitude", "longitude", "place_id", "busca",
]


def _linha(lead: Lead) -> dict:
    d = asdict(lead)
    d["prioridade"] = prioridade(lead)
    d["situacao_site"] = situacao_site(lead.site)
    d["whatsapp"] = whatsapp(lead.telefone)
    return d


def exportar(leads: list[Lead], base: str) -> list[str]:
    """Grava <base>.csv (separador ';', abre direto no Excel) e <base>.xlsx."""
    linhas = [_linha(l) for l in leads]
    caminhos = [f"{base}.csv"]
    with open(caminhos[0], "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=COLUNAS, delimiter=";", extrasaction="ignore")
        w.writeheader()
        w.writerows(linhas)

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return caminhos

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append(COLUNAS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    cores = {"A": "C6EFCE", "B": "FFEB9C", "C": "F2F2F2"}
    for d in linhas:
        ws.append([d.get(c) for c in COLUNAS])
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=cores[d["prioridade"]])
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    larguras = {"nome": 38, "categoria": 22, "endereco": 50, "site": 35, "link_maps": 30, "busca": 25}
    for i, c in enumerate(COLUNAS, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = larguras.get(c, 14)

    resumo = wb.create_sheet("Resumo")
    resumo.append(["Prioridade", "Significado", "Quantidade"])
    for p, txt in [("A", "Sem site e com telefone: melhor alvo"),
                   ("B", "Só rede social, ou sem site e sem telefone"),
                   ("C", "Já tem site: oferecer sistema/redesign")]:
        resumo.append([p, txt, sum(1 for d in linhas if d["prioridade"] == p)])
    resumo.append(["Total", "", len(linhas)])
    resumo.column_dimensions["B"].width = 45

    caminhos.append(f"{base}.xlsx")
    wb.save(caminhos[1])
    return caminhos
